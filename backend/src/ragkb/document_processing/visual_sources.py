"""Extract original embedded image bytes with real source locations; no size-based omission."""

from __future__ import annotations

import io
import json
import posixpath
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any

from pypdf import PdfReader

from ragkb.application.cancellation import check_cancelled
from ragkb.document_processing.docx_positions import paragraph_positions
from ragkb.domain.documents import SourceLocator


@dataclass(frozen=True)
class SourceImage:
    data: bytes
    locator: SourceLocator
    context: str = ""
    section_path: str = "root"
    caption: str = ""


def native_images(source: Path, kind: str, max_bytes: int) -> list[SourceImage]:
    images: list[SourceImage] = []
    if kind == "image":
        return [SourceImage(source.read_bytes(), SourceLocator(part=source.name))]
    if kind == "pptx":
        from pptx import Presentation

        presentation = Presentation(str(source))
        for number, slide in enumerate(presentation.slides, 1):
            title = slide.shapes.title.text.strip() if slide.shapes.title else ""
            section = f"第 {number} 张幻灯片" + (f" / {title}" if title else "")
            for blip in slide.element.xpath(".//a:blip"):
                identity = blip.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                )
                if identity and identity in slide.part.rels:
                    part = slide.part.related_part(identity)
                    images.append(
                        SourceImage(
                            part.blob,
                            SourceLocator(slide=number, part=str(part.partname).lstrip("/")),
                            title,
                            section,
                        )
                    )
    if kind == "xlsx":
        images.extend(_xlsx_images(source, max_bytes))
    if kind == "docx":
        from docx import Document

        document = Document(str(source))
        positions = paragraph_positions(document)
        for index, position in enumerate(positions):
            neighbours = [
                p
                for p in positions[max(0, index - 2) : index + 3]
                if p.section_path == position.section_path
                and p.element.getparent() is position.element.getparent()
            ]
            for blip in position.element.xpath(".//a:blip"):
                identity = blip.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                )
                part = document.part.related_parts.get(identity)
                if part is not None:
                    images.append(
                        SourceImage(
                            part.blob,
                            SourceLocator(
                                part=str(part.partname).lstrip("/"),
                                paragraph=position.number,
                            ),
                            "\n".join(p.text for p in neighbours if p.text)[:4000],
                            position.section_path,
                            "\n".join(p.text for p in neighbours if p.is_caption),
                        )
                    )
        # Headers, footers and unsupported drawing forms still retain their original media.
    if kind == "pdf":
        for number, page in enumerate(PdfReader(source).pages, 1):
            check_cancelled()
            for image in page.images:
                images.append(
                    # A page can contain several products/columns. Without an exact
                    # anchor its entire text is not this image's caption or context.
                    SourceImage(image.data, SourceLocator(page=number, part=image.name))
                )
    elif kind in {"docx", "pptx", "xlsx"}:
        prefix = {"docx": "word/media/", "pptx": "ppt/media/", "xlsx": "xl/media/"}[kind]
        with zipfile.ZipFile(source) as archive:
            for member in archive.infolist():
                check_cancelled()
                if member.filename.startswith(prefix) and not member.is_dir():
                    # Skip audio/video media, never pictures based on byte size or resolution.
                    if PurePosixPath(member.filename).suffix.lower() not in {
                        ".png",
                        ".jpg",
                        ".jpeg",
                        ".webp",
                        ".gif",
                        ".bmp",
                        ".tif",
                        ".tiff",
                        ".emf",
                        ".wmf",
                        ".svg",
                    }:
                        continue
                    if member.file_size > max_bytes:
                        raise ValueError("OCR_IMAGE_BYTES_LIMIT")
                    if any(image.locator.part == member.filename for image in images):
                        continue
                    images.append(
                        SourceImage(archive.read(member), SourceLocator(part=member.filename))
                    )
    return images


def _xlsx_images(source: Path, max_bytes: int) -> list[SourceImage]:
    """Read drawing anchors directly: shared image bytes can occur on different sheets."""
    images: list[SourceImage] = []
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    from openpyxl.utils.cell import get_column_letter

    with zipfile.ZipFile(source) as archive:

        def xml(part: str) -> Any:
            if archive.getinfo(part).file_size > max_bytes:
                raise ValueError("OCR_IMAGE_BYTES_LIMIT")
            return ET.fromstring(archive.read(part))  # noqa: S314 -- bounded OOXML part

        def related(part: str) -> dict[str, str]:
            path = PurePosixPath(part)
            relfile = str(path.parent / "_rels" / (path.name + ".rels"))
            if relfile not in archive.namelist():
                return {}
            return {
                item.attrib["Id"]: posixpath.normpath(str(path.parent / item.attrib["Target"]))
                if not item.attrib["Target"].startswith("/")
                else item.attrib["Target"].lstrip("/")
                for item in xml(relfile)
                if item.get("TargetMode") != "External"
            }

        sheets = related("xl/workbook.xml")
        for sheet in xml("xl/workbook.xml").findall(".//{*}sheet"):
            part = sheets.get(sheet.get(rel_ns + "id"))
            if not part:
                continue
            drawings = related(part)
            for drawing in xml(part).findall(".//{*}drawing"):
                drawing_part = drawings.get(drawing.get(rel_ns + "id"))
                if not drawing_part:
                    continue
                media = related(drawing_part)
                for anchor in xml(drawing_part):
                    start = anchor.find("{*}from")
                    row = int(start.findtext("{*}row", "0")) + 1 if start is not None else None
                    col = int(start.findtext("{*}col", "0")) + 1 if start is not None else None
                    for blip in anchor.findall(".//{*}blip"):
                        image_part = media.get(blip.get(rel_ns + "embed"))
                        if not image_part:
                            continue
                        if archive.getinfo(image_part).file_size > max_bytes:
                            raise ValueError("OCR_IMAGE_BYTES_LIMIT")
                        name = sheet.attrib["name"]
                        images.append(
                            SourceImage(
                                archive.read(image_part),
                                SourceLocator(
                                    sheet=name,
                                    row=row,
                                    cell_range=f"{get_column_letter(col)}{row}"
                                    if col and row
                                    else None,
                                    part=image_part,
                                ),
                                section_path=name,
                            )
                        )
    return images


def mineru_images(payload: bytes, max_bytes: int) -> list[SourceImage]:
    images: list[SourceImage] = []
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        paths = {entry.filename: entry for entry in archive.infolist()}
        content_names = [n for n in paths if n.endswith("content_list.json")]
        if len(content_names) != 1:
            raise ValueError("MINERU_VISUAL_CONTENT_AMBIGUOUS")
        name = content_names[0]
        base = PurePosixPath(name).parent
        content = json.loads(archive.read(name))
        for item in content:
            check_cancelled()
            if str(item.get("type", "")).casefold() not in {"image", "chart", "table"}:
                continue
            image_path = item.get("img_path") or item.get("image_path")
            if not image_path:
                continue
            path = PurePosixPath(str(image_path).replace("\\", "/"))
            if path.is_absolute() or ".." in path.parts:
                raise ValueError("MINERU_VISUAL_PATH_INVALID")
            matches = [n for n in paths if n in {str(path), str(base / path)}]
            if len(matches) != 1:
                raise ValueError("MINERU_VISUAL_FILE_MISSING")
            entry = paths[matches[0]]
            if entry.file_size > max_bytes:
                raise ValueError("OCR_IMAGE_BYTES_LIMIT")
            bbox: Any = item.get("bbox")
            locator = SourceLocator(
                page=int(item["page_idx"]) + 1, bbox=tuple(bbox) if bbox else None, part=str(path)
            )
            raw_caption = item.get("image_caption") or item.get("table_caption") or ""
            caption = (
                "\n".join(str(value) for value in raw_caption)
                if isinstance(raw_caption, list)
                else str(raw_caption)
            )[:4000]
            images.append(SourceImage(archive.read(entry), locator, caption, caption=caption))
    return images
