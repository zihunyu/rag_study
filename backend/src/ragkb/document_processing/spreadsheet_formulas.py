"""Bounded, deterministic arithmetic for spreadsheet formulas; no executable expressions."""

import re
from datetime import date, time
from decimal import ROUND_HALF_UP, Decimal, DecimalException
from typing import Any

from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import get_column_letter, range_boundaries


class FormulaUnavailable(ValueError):
    pass


def number(value: Any) -> Decimal:
    if value is None:
        return Decimal(0)
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        result = Decimal(str(value))
        if result.is_finite() and abs(result) <= Decimal("1e100"):
            return result
    raise FormulaUnavailable("引用值不是可计算的有限数值")


class FormulaEvaluator:
    def __init__(self, cells: dict[tuple[str, str], Any]) -> None:
        self.cells = cells
        self.sheets = {s for s, _ in cells}
        self.computed: dict[tuple[str, str], Decimal] = {}
        self.active: set[tuple[str, str]] = set()
        self.visits = 0

    def cell(self, sheet: str, address: str) -> Any:
        self.visits += 1
        if self.visits > 100000:
            raise FormulaUnavailable("公式计算预算已用完")
        key = (sheet, address.replace("$", "").upper())
        raw = self.cells.get(key)
        if isinstance(raw, (date, time)) or (
            isinstance(raw, str)
            and raw.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!"))
        ):
            raise FormulaUnavailable("引用了日期或错误单元格，需由表格程序复核")
        if not isinstance(raw, str) or not raw.startswith("="):
            return raw
        if key in self.computed:
            return self.computed[key]
        if key in self.active or len(self.active) >= 64:
            raise FormulaUnavailable("循环引用或依赖过深")
        self.active.add(key)
        try:
            result = number(Expression(raw, sheet, self).evaluate())
            self.computed[key] = result
            return result
        except (DecimalException, ZeroDivisionError, OverflowError, RecursionError) as error:
            raise FormulaUnavailable("公式计算失败") from error
        finally:
            self.active.remove(key)

    def reference(self, sheet: str, value: str) -> Any:
        if "!" in value:
            sheet, value = value.rsplit("!", 1)
            sheet = sheet.strip("'").replace("''", "'")
        if "[" in sheet or "]" in sheet or sheet not in self.sheets:
            raise FormulaUnavailable("外部或未知工作表引用")
        if not re.fullmatch(r"\$?[A-Za-z]{1,3}\$?[1-9]\d*(?::\$?[A-Za-z]{1,3}\$?[1-9]\d*)?", value):
            raise FormulaUnavailable("不支持的名称或范围引用")
        if ":" not in value:
            return self.cell(sheet, value)
        left, top, right, bottom = range_boundaries(value.replace("$", ""))
        if left is None or top is None or right is None or bottom is None:
            raise FormulaUnavailable("不完整的单元格范围")
        if right < left or bottom < top:
            raise FormulaUnavailable("倒置的单元格范围")
        if (right - left + 1) * (bottom - top + 1) > 10000:
            raise FormulaUnavailable("公式范围超过计算预算")
        return [
            self.cell(sheet, f"{get_column_letter(c)}{r}")
            for r in range(top, bottom + 1)
            for c in range(left, right + 1)
        ]


class Expression:
    def __init__(self, formula: str, sheet: str, evaluator: FormulaEvaluator) -> None:
        if len(formula) > 8192:
            raise FormulaUnavailable("公式过长")
        try:
            self.tokens = [t for t in Tokenizer(formula).items if t.type != "WHITE-SPACE"]
        except Exception as error:
            raise FormulaUnavailable("无法解析公式") from error
        if len(self.tokens) > 512:
            raise FormulaUnavailable("公式过于复杂")
        self.position = 0
        self.sheet, self.evaluator = sheet, evaluator

    def peek(self) -> str:
        return self.tokens[self.position].value if self.position < len(self.tokens) else ""

    def take(self, expected: str) -> None:
        if self.peek() != expected:
            raise FormulaUnavailable("不支持的公式结构")
        self.position += 1

    def evaluate(self) -> Any:
        value = self.add()
        if self.position != len(self.tokens):
            raise FormulaUnavailable("不支持的公式运算")
        return value

    def add(self) -> Any:
        value = self.multiply()
        while self.peek() in {"+", "-"}:
            op = self.peek()
            self.position += 1
            right = number(self.multiply())
            value = number(value) + right if op == "+" else number(value) - right
        return value

    def multiply(self) -> Any:
        value = self.atom()
        while self.peek() in {"*", "/"}:
            op = self.peek()
            self.position += 1
            right = number(self.atom())
            value = number(value) * right if op == "*" else number(value) / right
        return value

    def atom(self) -> Any:
        sign = Decimal(1)
        if self.peek() in {"+", "-"}:
            sign = Decimal(-1) if self.peek() == "-" else Decimal(1)
            self.position += 1
        if self.position >= len(self.tokens):
            raise FormulaUnavailable("公式不完整")
        token = self.tokens[self.position]
        self.position += 1
        if token.type == "FUNC" and token.subtype == "OPEN":
            args = []
            if self.peek() != ")":
                args.append(self.add())
                while self.peek() == ",":
                    self.position += 1
                    args.append(self.add())
            self.take(")")
            value = self.function(token.value[:-1].upper(), args)
        elif token.type == "PAREN" and token.subtype == "OPEN":
            value = self.add()
            self.take(")")
        elif token.type == "OPERAND" and token.subtype == "NUMBER":
            value = Decimal(token.value)
        elif token.type == "OPERAND" and token.subtype == "RANGE":
            value = self.evaluator.reference(self.sheet, token.value)
        else:
            raise FormulaUnavailable("公式含未支持的函数或数据类型")
        while self.peek() == "%":
            self.position += 1
            value = number(value) / 100
        return number(value) * sign if sign == -1 else value

    @staticmethod
    def function(name: str, args: list[Any]) -> Decimal:
        values = [v for arg in args for v in (arg if isinstance(arg, list) else [arg])]
        numeric = [
            number(v)
            for v in values
            if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)
        ]
        if name == "SUM":
            return sum(numeric, Decimal(0))
        if name == "COUNT":
            return Decimal(len(numeric))
        if name in {"MIN", "MAX"}:
            return (min if name == "MIN" else max)(numeric, default=Decimal(0))
        if name == "AVERAGE" and numeric:
            return sum(numeric, Decimal(0)) / len(numeric)
        if name == "ABS" and len(args) == 1:
            return abs(number(args[0]))
        if name == "ROUND" and len(args) == 2:
            digits = number(args[1])
            if digits == int(digits) and abs(digits) <= 20:
                return number(args[0]).quantize(Decimal(10) ** -int(digits), rounding=ROUND_HALF_UP)
        raise FormulaUnavailable("函数尚未支持或参数无效")
