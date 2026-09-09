"""Minimal answer reference presentation; management locators stay behind management APIs."""

from __future__ import annotations

from typing import Any, cast

from openpyxl.utils.cell import get_column_letter, range_boundaries


def _parent_cell_range(locator: dict[str, Any]) -> str | None:
    """Only merge witnessed, contiguous rows with the same columns and sheet."""
    spans = locator.get("source_spans")
    if not isinstance(spans, list) or not spans:
        return None
    bounds = []
    sheet = locator.get("sheet", locator.get("sheet_name"))
    for span in spans:
        child = span.get("locator") if isinstance(span, dict) else None
        if not isinstance(child, dict) or child.get("sheet", child.get("sheet_name")) != sheet:
            return None
        value = child.get("cell_range")
        if not isinstance(value, str):
            return None
        try:
            boundary = range_boundaries(value)
        except ValueError:
            return None
        if any(n is None for n in boundary):
            return None
        bounds.append(cast(tuple[int, int, int, int], boundary))
    ordered = sorted(bounds, key=lambda b: b[1])
    left, top, right, bottom = ordered[0]
    for a, start, b, end in ordered:
        if (a, b) != (left, right) or start > bottom + 1:
            return None
        bottom = max(bottom, end)
    return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"


def citation_locator(locator: dict[str, Any]) -> dict[str, Any]:
    projected = {
        key: value
        for key, value in locator.items()
        if key
        in {
            "page",
            "page_number",
            "page_start",
            "page_end",
            "paragraph",
            "paragraph_index",
            "start_line",
            "end_line",
            "section_path",
            "heading_path",
            "heading",
            "sheet",
            "sheet_name",
            "slide",
            "slide_number",
            "row_start",
            "row_end",
            "cell_range",
            "filename",
            "caption",
            "document_name",
            "visual_asset_ids",
            "source_type",
        }
    }
    if locator.get("is_parent") and "cell_range" in projected:
        # A parent inherits its first child's locator during chunking. Its
        # displayed text may cover many rows; never label it as just row 1.
        cell_range = _parent_cell_range(locator)
        if cell_range:
            projected["cell_range"] = cell_range
        else:
            projected.pop("cell_range")
    return projected


def cited_asset_ids(locator: dict[str, Any]) -> list[str]:
    ids = [i for i in locator.get("visual_asset_ids", []) if isinstance(i, str)]
    if "used_visual_fact_ids" not in locator:
        return ids  # Older runs attach only the cited chunk's original source images.
    receipt = locator.get("used_visual_fact_ids", [])
    if not isinstance(receipt, list):
        return []
    return [
        asset
        for asset in ids
        if any(isinstance(fact, str) and f":{asset}:" in fact for fact in receipt)
    ]


def reader_report(report: dict[str, Any]) -> dict[str, Any]:
    """An execution report must not become an uncited document/section listing."""
    value: dict[str, Any] = {
        k: v
        for k, v in report.items()
        if k
        in {
            "mode",
            "complete",
            "total_sections",
            "read_sections",
            "read_chunks",
            "checked_images",
            "answer_sections",
            "total_chunks",
        }
        and isinstance(v, (str, int, float, bool))
    }
    if isinstance(report.get("conditions"), dict):
        value["conditions"] = {
            k: v
            for k, v in report["conditions"].items()
            if k in {"checked", "covered", "missing", "not_applicable", "complete"}
            and isinstance(v, (int, bool))
        }
    failure = report.get("verification_failure")
    if isinstance(failure, dict):
        value["verification_failure"] = {
            k: v
            for k, v in failure.items()
            if (
                k
                in {
                    "batch_number",
                    "batch_count",
                    "completed_batches",
                    "condition_count",
                    "http_status",
                }
                and type(v) is int
            )
            or (
                k == "stage"
                and v
                in {
                    "claims_and_conflicts",
                    "conditions",
                    "condition_repair",
                    "citation_repair",
                    "answer_surface_repair",
                }
            )
            or (
                k == "provider_code"
                and v
                in {
                    "MODEL_ACCOUNT_QUOTA_WAIT_TIMEOUT",
                    "MODEL_PROVIDER_DEADLINE_EXCEEDED",
                    "MODEL_PROVIDER_TIMEOUT",
                    "MODEL_PROVIDER_RATE_LIMITED",
                }
            )
        }
    execution_failure = report.get("execution_failure")
    if isinstance(execution_failure, dict):
        from ragkb.application.conversation_diagnostics import PUBLIC_CONVERSATION_CODES

        value["execution_failure"] = {
            k: v
            for k, v in execution_failure.items()
            if (
                k == "stage"
                and isinstance(v, str)
                and v in {"conversation_context", "knowledge_qa"}
            )
            or (k == "code" and isinstance(v, str) and v in PUBLIC_CONVERSATION_CODES)
            or (k == "http_status" and type(v) is int and 400 <= v <= 599)
            or (
                k == "request_id"
                and isinstance(v, str)
                and len(v) == 36
                and all(c in "0123456789abcdef-" for c in v)
            )
        }
    performance = report.get("performance")
    if isinstance(performance, dict) and type(performance.get("elapsed_seconds")) in {int, float}:
        value["performance"] = {
            "elapsed_seconds": performance["elapsed_seconds"],
            "events": [
                {
                    k: v
                    for k, v in event.items()
                    if k
                    in {
                        "kind",
                        "name",
                        "seconds",
                        "status",
                        "batch_number",
                        "sent",
                        "queue_seconds",
                        "network_seconds",
                        "cache",
                        "outcome",
                    }
                    and type(v) in {str, int, float, bool}
                }
                for event in performance.get("events", [])[:512]
                if isinstance(event, dict)
            ],
        }
    value["image_checks"] = [
        {"status": c["status"]}
        for c in report.get("image_checks", [])
        if isinstance(c, dict)
        and c.get("status")
        in {
            "supported",
            "uncertain",
            "conflict",
            "verification_failed",
            "irrelevant",
            "budget_exceeded",
        }
    ]
    if report.get("gaps"):
        value["gaps"] = ["本轮尚有未覆盖的资料或条件，请缩小问题范围后重试。"]
    return value
