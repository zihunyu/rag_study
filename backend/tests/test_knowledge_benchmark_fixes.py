import sqlite3
from types import SimpleNamespace as NS

import openpyxl
import pytest
from ragkb.application.reading_scope import ReadingOptions, reading_scope
from ragkb.document_processing.chunking import TokenAwareChunker
from ragkb.document_processing.office_parsers import SpreadsheetParser
from ragkb.domain.question_coverage import question_aspects
from ragkb.infrastructure.overview import OverviewReader


def overview_fixture(*, cache_error="", search=None):
    docs = [
        {"document_id": "base", "version_id": "base", "filename": "星云设备.md"},
        {"document_id": "materials", "version_id": "materials", "filename": "星云设备申请材料.md"},
    ]
    sources = {
        d["document_id"]: NS(
            chunk_id=d["document_id"],
            document_id=d["document_id"],
            retrieval_text="保修三年。" if d["document_id"] == "base" else "申请需身份证与发票。",
            display_text="",
            locator={"section_path": "正文"},
            valid_from_epoch=0,
            valid_to_epoch=0,
            permission_revision=1,
        )
        for d in docs
    }

    class Repository:
        def list_documents_page(self, *args, **kwargs):
            return NS(items=docs, next_key=None)

        def list_chunks_page(self, version, **kwargs):
            return NS(items=[{"chunk_id": version}], next_key=None)

    class Ledger:
        def get(self, *args):
            return {}

        def cache_get(self, *args):
            if cache_error == "read":
                raise sqlite3.OperationalError("locked")
            return None

        def cache_put(self, *args):
            if cache_error in {"read", "write"}:
                raise sqlite3.OperationalError("disk unavailable")

    calls = []

    class Reader:
        revision = "fixture"

        def read(self, question, batch):
            calls.append(1)
            return {
                "quotes": [{"chunk_id": e.chunk_id, "quote": e.text} for e in batch],
                "gaps": [],
            }

    service = OverviewReader(
        Repository(),
        NS(authorize_chunks=lambda ids, ctx: {i: sources[i] for i in ids}),
        NS(ledger=Ledger(), list_assets=lambda v: []),
        NS(
            overview_max_chunks=100,
            overview_evidence_tokens=1000,
            overview_max_images=0,
            ocr_generation_cache_ttl_seconds=600,
        ),
        Reader(),
        scope_search=search,
    )
    return service, NS(space_ids=("space",), tenant_id="tenant"), calls


@pytest.mark.parametrize("failure", ["read", "write"])
def test_chapter_cache_failure_returns_computed_sources_and_reuses_them(failure):
    reader, context, calls = overview_fixture(cache_error=failure)
    reader.settings.overview_evidence_tokens = 6
    with reading_scope(ReadingOptions(document_ids=("base", "materials"))):
        first, report = reader.read("总结全文", context)
        count = len(calls)
        second, again = reader.read("总结全文", context)
    assert count == 2 and len(calls) == count
    assert first == second and report == again
    assert first
    assert not any("生成失败" in gap for gap in report["gaps"])


def test_title_seed_always_finds_companions_and_does_not_claim_confirmed_scope():
    searched = []
    reader, context, _ = overview_fixture(search=lambda *args: searched.append(1) or ("materials",))
    sources, report = reader.read("总结星云设备的保修及申请材料", context)
    assert searched == [1]
    assert {e.document_id for e in sources} == {"base", "materials"}
    assert report["selected_documents_read_complete"]
    assert not report["scope_complete"] and not report["complete"]
    with reading_scope(ReadingOptions(document_ids=("base", "materials"))):
        _, confirmed = reader.read("总结全文", context)
    assert confirmed["scope_complete"] and confirmed["complete"]
    _, whole = reader.read("总结整个知识库", context)
    assert whole["scope_resolution"] == "whole_space"
    assert whole["scope_confirmed"] and whole["complete"]


@pytest.mark.parametrize(
    ("question", "expected"),
    [
        ("A款支持哪些接口，B款呢？", ("A款支持哪些接口", "B款支持哪些接口")),
        ("A款的保修多久，B款呢，C款呢？", ("A款的保修多久", "B款的保修多久", "C款的保修多久")),
        (
            "介绍星云设备的优点、缺点和适用人群",
            ("星云设备的优点", "星云设备的缺点", "星云设备的适用人群"),
        ),
        (
            "比较A款和B款的价格和保修期限",
            ("A款的价格", "A款的保修期限", "B款的价格", "B款的保修期限"),
        ),
        ("请介绍“温度和压力”的价格和精度", ("“温度和压力”的价格", "“温度和压力”的精度")),
    ],
)
def test_explicit_requirements_include_ellipsis_and_each_matrix_cell(question, expected):
    assert question_aspects(question) == expected


def test_xlsx_title_header_formulas_and_chunk_context(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["2026销售统计（元）", None, None, None])
    sheet.merge_cells("A1:D1")
    sheet.append(["产品", "数量", "单价", "金额"])
    sheet.append(["A款", 2, 100, "=B3*C3"])
    sheet.append(["B款", 3, 120, "=B4*C4"])
    sheet.append(["合计", None, None, "=SUM(D3:D4)"])
    file = tmp_path / "sales.xlsx"
    workbook.save(file)
    parsed = SpreadsheetParser().parse(file, "version")
    assert parsed.nodes[2].metadata["table_header"] == "产品 | 数量 | 单价 | 金额"
    assert parsed.nodes[2].metadata["table_title"] == "2026销售统计（元）"
    assert parsed.nodes[2].metadata["formulas"][0]["value"] == "200"
    assert parsed.nodes[-1].metadata["formulas"][0]["value"] == "560"
    chunks = TokenAwareChunker().chunk(parsed, tenant_id="tenant").chunks
    matching = [c for c in chunks if "=B3*C3" in c.original_text]
    assert matching and "200" in matching[0].retrieval_text
    assert "2026销售统计（元）" in matching[0].retrieval_text
    assert "产品 | 数量 | 单价 | 金额" in matching[0].retrieval_text


@pytest.mark.parametrize(
    "formula", ["=A2", "=VLOOKUP(1,A1:B2,2,0)", "='[secret.xlsx]Sheet'!A1", "=1/0"]
)
def test_uncomputable_formulas_are_explicitly_unknown(tmp_path, formula):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["金额"])
    sheet.append([formula])
    path = tmp_path / "unknown.xlsx"
    workbook.save(path)
    node = SpreadsheetParser().parse(path, "v").nodes[1]
    assert node.metadata["formulas"][0]["status"] == "unresolved"
    assert "不能作为已知数值使用" in node.original_text


def test_multiple_followups_and_quoted_entity_names():
    assert question_aspects("A款支持哪些接口，那么B款和C款呢？") == (
        "A款支持哪些接口",
        "B款支持哪些接口",
        "C款支持哪些接口",
    )
    assert question_aspects("A款的价格，‘B与C’呢？") == ("A款的价格", "‘B与C’的价格")


def test_multilevel_merged_headers_keep_units_and_coordinates(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["产品", "销售（元）", None, "成本（元）"])
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:C1")
    sheet.merge_cells("D1:D2")
    sheet.append([None, "数量", "金额", None])
    sheet.append(["A款", 2, 200, 100])
    sheet.append([])
    sheet.append(["名称", "数量"])
    sheet.append(["B款", 3])
    path = tmp_path / "merged.xlsx"
    book.save(path)
    result = SpreadsheetParser().parse(path, "v")
    data = next(n for n in result.nodes if n.metadata["row"] == 3)
    assert data.metadata["column_labels"] == [
        "产品",
        "销售（元） / 数量",
        "销售（元） / 金额",
        "成本（元）",
    ]
    assert data.metadata["header_rows"] == [1, 2]
    assert data.metadata["source_spans"][1]["locator"]["cell_range"] == "A1:D2"
    assert result.nodes[-1].metadata["table_id"] != data.metadata["table_id"]


def test_formula_cross_sheet_arithmetic_rounding_and_errors():
    from decimal import Decimal

    from ragkb.document_processing.spreadsheet_formulas import FormulaEvaluator, FormulaUnavailable

    cells = {
        ("价格", "A1"): 19.95,
        ("订单", "B1"): 3,
        ("订单", "C1"): "=ROUND('价格'!$A$1*B1*(1-10%),2)",
        ("订单", "C2"): "=SUM(C1:C1)",
        ("订单", "A2"): "#DIV/0!",
        ("订单", "D2"): "=SUM(A2:A2)",
    }
    evaluator = FormulaEvaluator(cells)
    assert evaluator.cell("订单", "C1") == Decimal("53.87")
    assert evaluator.cell("订单", "C2") == Decimal("53.87")
    with pytest.raises(FormulaUnavailable):
        evaluator.cell("订单", "D2")


def test_chapter_memory_copy_expires_is_bounded_and_does_not_alias(monkeypatch):
    from ragkb.infrastructure import chapter_cache
    from ragkb.infrastructure.chapter_cache import ChapterCache, valid_chapter

    now = [0]
    monkeypatch.setattr(chapter_cache.time, "monotonic", lambda: now[0])

    class Broken:
        def cache_get(self, *args):
            return None

        def cache_put(self, *args):
            raise sqlite3.OperationalError("locked")

    cache = ChapterCache(Broken(), max_bytes=100)
    payload = {"quotes": [{"chunk_id": "c", "quote": "证据"}], "gaps": []}
    cache.put("t", "k", payload, 10)
    payload["quotes"].clear()
    assert cache.get("t", "k")["quotes"]
    cache.put("t", "big", {"huge": "字" * 100}, 10)
    assert cache.get("t", "big") is None and cache.size <= 100
    now[0] = 11
    assert cache.get("t", "k") is None
    assert not valid_chapter({"quotes": [], "gaps": []}, {})
    assert not valid_chapter({"quotes": [{"chunk_id": "c", "quote": "伪造"}]}, {"c": "证据"})
