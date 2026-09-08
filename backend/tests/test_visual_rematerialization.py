"""A new draft can rebuild proven visual facts without rewriting prior human history."""

from __future__ import annotations

import copy
import hashlib
import importlib.util
import io
import json
from argparse import Namespace
from dataclasses import replace
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from ragkb.api.app import create_app
from ragkb.application.worker import LocalIngestionWorker
from ragkb.document_processing.parsers import ParserRouter
from ragkb.document_processing.visual_parser import VisualDocumentParser
from ragkb.domain.documents import CanonicalDocument, CanonicalNode, NodeType, SourceLocator
from ragkb.infrastructure.visual_assets import VisualAssetStore
from ragkb.infrastructure.visual_rematerialization import (
    REVISION,
    apply_snapshot,
    build_snapshot,
    digest,
    materialized_visual_text,
    rematerialize_canonical,
)
from test_graph_identity_and_queries import extraction, fixture
from test_visual_pipeline import png
from test_workspace_redesign import publish
from test_workspace_redesign import workspace as workspace_fixture

workspace = workspace_fixture


@pytest.fixture
def reviewed_source(workspace):
    client, runtime, _, space = workspace
    store = VisualAssetStore(runtime.storage)
    doc = Document()
    doc.add_paragraph("原生正文保持不变")
    doc.add_picture(io.BytesIO(png()))
    output = io.BytesIO()
    doc.save(output)
    data = output.getvalue()
    created = client.post(
        f"/api/spaces/{space}/upload-sessions",
        headers={"Idempotency-Key": "legacy-reviewed"},
        json={
            "filename": "reviewed.docx",
            "expected_size": len(data),
            "expected_sha256": hashlib.sha256(data).hexdigest(),
            "declared_mime": "application/octet-stream",
        },
    )
    assert created.status_code == 201, created.text
    uploaded = client.put(
        created.json()["upload_path"], headers={"If-Match": created.headers["etag"]}, content=data
    )
    response = client.post(
        f"/api/upload-sessions/{created.json()['upload_session_id']}:complete",
        headers={
            "If-Match": uploaded.headers["etag"],
            "Idempotency-Key": "legacy-reviewed-complete",
        },
    )
    assert response.status_code == 202, response.text
    version = response.json()["document_version_id"]

    class LegacyParser:
        revision = "legacy-reviewed-visuals"

        def parse(self, source, identity):
            asset = store.save_image(
                identity, png(), {"part": "word/media/image1.png", "paragraph": 2}
            )
            asset.update(
                status="verified",
                origin="human_review",
                revision="old-ocr-revision",
                extraction=extraction(fixture()).model_dump(),
                issues=[],
                analyzed_at=1234,
                history=[{"actor": "reviewer", "reason": "已逐条对照原图核实", "at": 1234}],
                section_path="正文",
                context="原生正文保持不变",
            )
            store.write_manifest(identity, [asset])
            return CanonicalDocument(
                identity,
                "zh",
                "docx",
                (
                    CanonicalNode(
                        "native",
                        None,
                        NodeType.PARAGRAPH,
                        "原生正文保持不变",
                        "原生正文保持不变",
                        SourceLocator(paragraph=1),
                    ),
                    CanonicalNode(
                        "old-image",
                        None,
                        NodeType.IMAGE,
                        "网关 指向 ；服务请求响应",
                        "网关 指向 ；服务请求响应",
                        SourceLocator(paragraph=2, part="word/media/image1.png"),
                        {"visual_asset_ids": [asset["id"]], "section_path": "正文"},
                    ),
                ),
                self.revision,
                "native-v1",
                hashlib.sha256(source.read_bytes()).hexdigest(),
                media_refs=(asset,),
            )

    worker = LocalIngestionWorker(
        runtime.queue,
        runtime.repository,
        runtime.storage,
        ParserRouter({"docx": LegacyParser()}),
        "legacy-worker",
        chunker=runtime.chunker,
        indexing_sink=runtime.indexing_sink,
    )
    assert worker.run_once()
    publish(client, version)
    configured = replace(
        runtime, settings=runtime.settings.model_copy(update={"ocr_enabled": True})
    )

    def bind(session_id, new_version):
        plan = store.ledger.get("session_plan", session_id)
        if plan:
            store.ledger.put("version_plan", new_version, plan, immutable=True)

    runtime.uploads.before_enqueue = bind
    with TestClient(create_app(configured)) as updated:
        yield updated, configured, store, response.json()


def command(client, version):
    response = client.get(f"/api/document-versions/{version}/visual-rematerialization")
    assert response.status_code == 200, response.text
    report = response.json()
    return (
        report,
        {"If-Match": str(report["row_version"]), "Idempotency-Key": "materialize-once"},
        {
            "input_fingerprint": report["input_fingerprint"],
            "backup_receipt_sha256": "a" * 64,
        },
    )


def test_dry_run_is_read_only_and_materialization_creates_idempotent_unpublished_version(
    reviewed_source,
):
    client, runtime, store, done = reviewed_source
    version, document = done["document_version_id"], done["document_id"]
    before = copy.deepcopy(store.list_assets(version))
    report, headers, body = command(client, version)
    assert len(runtime.repository.get_versions(document)) == 1
    assert "未命名区域（包含：NACOS、数据库）" in report["preview"][0]["text"]
    assert "指向 ；" not in report["preview"][0]["text"]
    result = client.post(
        f"/api/document-versions/{version}:rematerialize-visuals", headers=headers, json=body
    )
    assert result.status_code == 202, result.text
    new_version = result.json()["document_version_id"]
    assert new_version != version
    replay = client.post(
        f"/api/document-versions/{version}:rematerialize-visuals", headers=headers, json=body
    )
    assert replay.status_code == 202 and replay.json() == result.json()
    assert len(runtime.repository.get_versions(document)) == 2

    class Never:
        revision = "never-called"

        def parse(self, *args, **kwargs):
            pytest.fail("Rematerialization reran a document parser")

        def analyze(self, *args, **kwargs):
            pytest.fail("Rematerialization reran OCR")

    parser = VisualDocumentParser(Never(), "docx", Never(), store, runtime.settings)
    worker = LocalIngestionWorker(
        runtime.queue,
        runtime.repository,
        runtime.storage,
        ParserRouter({"docx": parser}),
        "rematerialization-worker",
        chunker=runtime.chunker,
        indexing_sink=runtime.indexing_sink,
    )
    assert worker.run_once()
    assert runtime.repository.get_version(new_version)["processing_state"] == "VALIDATED"
    assert runtime.repository.get_version(new_version)["publication_state"] == "DRAFT"
    runtime.lifecycle_store.reload()
    assert runtime.lifecycle_store.documents[document].active_version_id == version
    new_asset = store.list_assets(new_version)[0]
    assert new_asset["history"] == before[0]["history"]
    assert new_asset["analyzed_at"] == before[0]["analyzed_at"]
    assert new_asset["revision"] == before[0]["revision"]
    assert new_asset["materialization_revision"] == REVISION
    assert store.list_assets(version) == before
    text = "\n".join(c["text"] for c in runtime.repository.list_chunks(new_version, preview=True))
    assert "原生正文保持不变" in text and "未命名区域" in text and "指向 ；" not in text
    assert any(
        "指向 ；" in c["text"] for c in runtime.repository.list_chunks(version, preview=True)
    )


@pytest.mark.parametrize(
    "change", ["later_version", "revoked", "deleted", "snapshot", "origin", "receipt", "pixels"]
)
def test_changed_or_unproven_source_cannot_be_rebuilt(reviewed_source, change):
    client, runtime, store, done = reviewed_source
    version, document = done["document_version_id"], done["document_id"]
    _, headers, body = command(client, version)
    if change == "later_version":
        first = client.post(
            f"/api/document-versions/{version}:rematerialize-visuals", headers=headers, json=body
        )
        assert first.status_code == 202
        headers = {**headers, "Idempotency-Key": "different-command"}
    elif change == "revoked":
        assert (
            client.post(
                f"/api/documents/{document}:revoke",
                headers={"Idempotency-Key": "revoke-before-rebuild"},
            ).status_code
            == 200
        )
    elif change == "deleted":
        assert client.delete(
            f"/api/documents/{document}", headers={"Idempotency-Key": "delete-before-rebuild"}
        ).status_code in {200, 202}
    elif change == "snapshot":
        body["input_fingerprint"] = "b" * 64
    else:
        assets = store.list_assets(version)
        if change == "origin":
            assets[0]["origin"] = "model"
        elif change == "receipt":
            assets[0]["history"] = []
        else:
            runtime.storage.write_bytes("artifacts", assets[0]["storage_key"], b"changed-image")
        store.write_manifest(version, assets)
    count = len(runtime.repository.get_versions(document))
    result = client.post(
        f"/api/document-versions/{version}:rematerialize-visuals", headers=headers, json=body
    )
    assert result.status_code in {404, 409}, result.text
    assert len(runtime.repository.get_versions(document)) == count


def test_worker_rejects_modified_plan_and_preserves_excluded_graph_relations(
    reviewed_source, tmp_path
):
    _, runtime, store, done = reviewed_source
    snapshot = build_snapshot(runtime, runtime.repository.get_version(done["document_version_id"]))
    original = runtime.storage.path_for("original", snapshot["version"]["original_key"])
    plan = {"rematerialization": {"snapshot": snapshot, "fingerprint": "0" * 64}}
    with pytest.raises(ValueError, match="PLAN_INTEGRITY_INVALID"):
        apply_snapshot(store, original, "new", plan)
    obsolete = copy.deepcopy(snapshot)
    obsolete["generator_revision"] = "reviewed-graph-materialization-v1"
    with pytest.raises(ValueError, match="PLAN_INTEGRITY_INVALID"):
        apply_snapshot(
            store,
            original,
            "new",
            {"rematerialization": {"snapshot": obsolete, "fingerprint": digest(obsolete)}},
        )
    graph = snapshot["assets"][0]["extraction"]["graphs"][0]
    graph["edges"][0]["review_status"] = "excluded"
    snapshot["assets"][0]["extraction"]["body_text"] = "网关总能到达生产区"
    rebuilt = rematerialize_canonical(snapshot, "new")
    text = rebuilt.nodes[1].display_text
    assert "网关总能到达生产区" not in text and "服务请求响应" not in text


def test_old_derived_text_is_replaced_but_independent_human_conditions_survive():
    reviewed = extraction(fixture())
    reviewed.body_text = "仅在城区保修期内提供服务。"
    reviewed.transcription = (
        "网关\nNACOS\n数据库\n仅在城区保修期内提供服务。\n网关 双向连接 ；服务请求响应"
    )
    before = reviewed.model_dump()
    text = materialized_visual_text(reviewed)
    assert "双向连接 ；" not in text
    assert "网关 双向连接 生产区 / 未命名区域" in text
    assert text.count("仅在城区保修期内提供服务。") == 1
    assert reviewed.model_dump() == before
    reviewed.body_text = "另一个未能对应的对象 指向 ；需要人工核对"
    with pytest.raises(ValueError, match="UNBOUND_HUMAN_PROSE_REQUIRES_REVIEW"):
        materialized_visual_text(reviewed)


def test_cli_saves_verified_immutable_backups_before_apply_and_can_replay_saved_plan(
    reviewed_source, tmp_path
):
    client, _, _, done = reviewed_source
    version = done["document_version_id"]
    path = Path(__file__).parents[2] / "scripts/rematerialize_visuals.py"
    spec = importlib.util.spec_from_file_location("rematerialize_cli", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    args = Namespace(
        version_id=version,
        document_id=done["document_id"],
        plan=None,
        backup_dir=str(tmp_path / "backup"),
        apply=False,
    )
    report = client.get(f"/api/document-versions/{version}/visual-rematerialization").json()

    def transport(request):
        response = client.request(
            request.method, request.url.path, headers=dict(request.headers), content=request.content
        )
        if request.method == "POST":
            assert list((tmp_path / "backup").rglob("receipt.json"))
            assert list((tmp_path / "backup").rglob("original.bin"))
        return httpx.Response(
            response.status_code, content=response.content, headers=response.headers
        )

    with httpx.Client(
        base_url="http://local.test", transport=httpx.MockTransport(transport)
    ) as bridge:
        dry = module.run(args, bridge)
        assert dry["mode"] == "dry-run"
        folder = Path(dry["backup_directory"])
        assert (
            json.loads((folder / "diagnostic.json").read_text(encoding="utf-8"))[
                "input_fingerprint"
            ]
            == report["input_fingerprint"]
        )
        args.apply, args.plan = True, dry["plan"]
        first = module.run(args, bridge)
        second = module.run(args, bridge)
        assert first["result"] == second["result"] and first["result"]["creates_draft_only"]
        (folder / "original.bin").write_bytes(b"tampered backup")
        with pytest.raises(ValueError, match="BACKUP_ALREADY_EXISTS"):
            module.run(args, bridge)
