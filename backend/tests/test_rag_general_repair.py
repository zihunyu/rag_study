from __future__ import annotations

import json

import pytest
from docx import Document
from openpyxl import Workbook
from ragkb.adapters.evidence_selection import ModelEvidenceSelector
from ragkb.application.evidence import SearchBackedEvidenceProvider
from ragkb.contracts.rag import EvidenceSelection
from ragkb.document_processing.chunking import ChunkingConfig, TokenAwareChunker
from ragkb.document_processing.parsers import ParserRouter
from ragkb.domain.errors import InvalidProviderResponse
from ragkb.domain.numeric_facts import check_numeric_facts
from ragkb.domain.rag import Evidence, QuestionAssessment, QuestionDisposition
from ragkb.domain.retrieval import SearchHit, SearchResult, SearchSource
from test_model_http_adapters import _MockTransport, _settings


@pytest.mark.parametrize("suffix", [".md", ".csv", ".xlsx", ".docx", ".html"])
def test_unrelated_formats_preserve_table_meaning_and_locations(tmp_path, suffix):
    path = tmp_path / ("海岚设备规格" + suffix)
    rows = [["型号", "环境", "额定功率(W)"], ["Q7", "常温", "420"], ["Q7", "低温", "390"]]
    if suffix == ".md":
        path.write_text(
            "# 海岚设备\n\n## 规格\n\n"
            + "\n".join(["|".join(rows[0]), "---|---|---", *["|".join(r) for r in rows[1:]]]),
            encoding="utf-8",
        )
    elif suffix == ".csv":
        path.write_text("\n".join(",".join(r) for r in rows), encoding="utf-8")
    elif suffix == ".xlsx":
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.save(path)
    elif suffix == ".docx":
        doc = Document()
        doc.add_heading("海岚设备", 1)
        table = doc.add_table(rows=3, cols=3)
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                table.cell(i, j).text = value
        doc.save(path)
    else:
        path.write_text(
            "<h1>海岚设备</h1><table>"
            + "".join("<tr>" + "".join(f"<td>{v}</td>" for v in row) + "</tr>" for row in rows)
            + "</table>",
            encoding="utf-8",
        )
    parsed = (
        ParserRouter()
        .route(
            {".md": "markdown", ".csv": "csv", ".xlsx": "xlsx", ".docx": "docx", ".html": "html"}[
                suffix
            ]
        )
        .parse(path, "version")
    )
    result = TokenAwareChunker(
        ChunkingConfig(target_tokens=35, overlap_tokens=3, min_tokens=2, max_tokens=50)
    ).chunk(parsed, tenant_id="tenant")
    values = [c for c in result.chunks if "420" in c.display_text]
    assert values and all("额定功率(W)" in c.retrieval_text for c in values)
    assert all("常温" in c.display_text for c in values)
    assert all(c.locator.to_dict() for c in result.chunks)
    assert all(p.metadata["source_spans"] for p in result.parent_chunks)
    assert not any(set(c.display_text) <= set("-|:\n ") for c in result.chunks)


def test_markdown_hierarchy_steps_and_real_offsets(tmp_path):
    path = tmp_path / "维修手册.md"
    raw = (
        "# 设备A\n\n## 关机\n\n1. 保存记录。\n2. 关闭电源。\n\n"
        "## 清洁\n\n### 前提\n\n必须先断电。\n"
    )
    path.write_text(raw, encoding="utf-8")
    doc = ParserRouter().route("markdown").parse(path, "v")
    result = TokenAwareChunker().chunk(doc, tenant_id="t")
    for c in result.chunks:
        left, right = c.locator.char_range
        assert raw[left:right] == c.display_text
    assert any("1. 保存记录。\n2. 关闭电源。" in c.display_text for c in result.chunks)
    assert any(c.metadata["section_path"] == "设备A / 清洁 / 前提" for c in result.chunks)


def source(identity, text):
    return SearchSource(identity, "doc", "v", text, text, {"page": 1}, 0, 0, 1, True)


class Search:
    revision = "test"

    def __init__(self):
        self.calls = []

    def search(self, query, context, **kwargs):
        self.calls.append((query, context))
        pool = (source("intro", "设备Q7的介绍"), source("power", "Q7 额定功率 420 W"))
        if len(self.calls) > 1:
            pool += (source("temperature", "Q7 最低工作温度 -10摄氏度"),)
        hit = SearchHit(
            "intro",
            "doc",
            "v",
            pool[0].display_text,
            {"page": 1},
            1,
            1,
            ("dense",),
            display_text=pool[0].display_text,
            retrieval_text=pool[0].display_text,
        )
        return SearchResult((hit,), 1, review_sources=pool)

    def expand_parents(self, sources, context):
        return ()


def provider(search, selector):
    return SearchBackedEvidenceProvider(
        search,
        space_id="space",
        active_generation_id="g",
        active_permission_revision=lambda: 1,
        required_security_watermark=lambda: 1,
        prompt_revision="p",
        model_revision="m",
        final_evidence_count=8,
        evidence_selector=selector,
    )


def test_selected_review_evidence_is_promoted_with_bounded_scoped_supplement():
    class Selector:
        revision = "test"

        def select(self, question, evidence):
            selected = tuple(e.evidence_id for e in evidence if "420" in e.text or "-10" in e.text)
            return EvidenceSelection(selected, "partial", ("Q7 最低工作温度",))

    search = Search()
    package = provider(search, Selector()).build_package(
        "Q7功率及工作温度", "t", "u", space_id="specific", subject_scope_tokens=("scope",)
    )
    assert len(search.calls) == 2  # A still-partial second selection must not loop.
    assert all(
        c.space_ids == ("specific",) and c.tenant_id == "t" and c.subject_scope_tokens == ("scope",)
        for _, c in search.calls
    )
    assert any("420" in e.text for e in package.generation_evidence)
    assert any("-10" in e.text for e in package.generation_evidence)
    assert not any("介绍" in e.text for e in package.generation_evidence)
    assert len(package.retrieval_queries) == 2
    assert package.coverage == "partial"


def test_noun_lookup_uses_sources_but_ambiguous_entities_request_clarification():
    class Assessor:
        revision = "test"

        def assess(self, question):
            return QuestionAssessment(
                QuestionDisposition.NEEDS_CLARIFICATION, "missing_context", ("subject",)
            )

    class Selector:
        revision = "test"

        def select(self, question, evidence):
            return EvidenceSelection((), "ambiguous", (), "你指的是Q7设备还是Q7软件？")

    search = Search()
    p = provider(search, Selector())
    p.question_assessor = Assessor()
    package = p.build_package("Q7", "t", "u")
    assert len(search.calls) == 1 and package.disposition == QuestionDisposition.NEEDS_CLARIFICATION
    assert "设备" in package.clarification_question
    p.build_package("它", "t", "u")
    assert len(search.calls) == 1


@pytest.mark.parametrize(
    "patch",
    [
        {"source_ids": ["E99"]},
        {"source_ids": ["E1", "E1"]},
        {"queries": ["q"] * 3},
        {"coverage": "anything"},
        {"coverage": "ambiguous", "clarification": None},
    ],
)
def test_selector_rejects_fabricated_ids_and_unbounded_plans(tmp_path, patch):
    settings, _ = _settings(tmp_path)
    value = {
        "source_ids": ["E1"],
        "coverage": "sufficient",
        "queries": [],
        "clarification": None,
        **patch,
    }
    transport = _MockTransport({"choices": [{"message": {"content": json.dumps(value)}}]})
    evidence = Evidence("E1", "c", "d", "v", "Q7 420W", {"page": 1}, 0, 0, 0, 1, True, True)
    with pytest.raises(InvalidProviderResponse):
        ModelEvidenceSelector(settings, transport).select("Q7功率", (evidence,))


def test_unlabelled_coordinate_values_require_semantics_without_false_contradiction():
    text = "| 地点 | 坐标 | 说明 |\n| 实验室 | 142.9, -68.9, 594.6 | 测量记录 |"
    assert check_numeric_facts("坐标是142.9, -68.9, 594.6。", (text,)) == "uncertain"
    # Swapping coordinates must never be accepted by literal value presence alone.
    assert check_numeric_facts("坐标是594.6, -68.9, 142.9。", (text,)) != "supported"
    assert check_numeric_facts("高度为999米。", ("高度为100米。",)) == "mismatch"
