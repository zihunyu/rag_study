"""Reviewing one image must never re-run OCR or approve its unresolved siblings."""

from __future__ import annotations

import copy
from dataclasses import replace

import pytest
from ragkb.adapters.local_storage import LocalFileStorage
from ragkb.document_processing.visual_processing import VisualProcessor
from ragkb.document_processing.visual_sources import SourceImage
from ragkb.domain.documents import SourceLocator
from ragkb.domain.visual_review import review_issues
from ragkb.infrastructure.visual_assets import VisualAssetStore
from test_visual_pipeline import EXTRACTION, png, settings
from test_visual_workspace import CountingAnalyzer


@pytest.fixture
def revision_workspace(tmp_path):
    store = VisualAssetStore(LocalFileStorage(tmp_path / "store"))
    analyzer = CountingAnalyzer()
    processor = VisualProcessor(store, analyzer, settings())
    pictures = [
        SourceImage(png(), SourceLocator(part=f"page-{i}.png"), f"page {i}") for i in range(6)
    ]
    original = processor.process("v1", pictures, {}, "doc")
    for index, (_, asset) in enumerate(original[:5]):
        asset.update(
            status="needs_review",
            stage="needs_review",
            issues=[f"第 {index + 1} 张原始疑点"],
            audit=[{"stage": "original_read", "receipt": {"request": index}}],
        )
        store.ledger.upsert_asset("v1", asset)
    analyzer.calls = 0
    return store, analyzer, processor, pictures


def test_sequential_single_image_reviews_preserve_unreviewed_snapshots_without_calls(
    revision_workspace,
):
    store, analyzer, processor, pictures = revision_workspace
    base = {a["id"]: a for a in store.list_assets("v1")}
    ids = list(base)
    # Even new model/local-check revisions and old cache timestamps must not
    # convert an immutable manual review into a whole-document OCR refresh.
    analyzer.revision = "new-model-revision"
    processor.settings = processor.settings.model_copy(
        update={"ocr_cross_version_cache_enabled": False, "ocr_local_check_enabled": False}
    )
    for number in range(1, 4):
        previous, target = f"v{number}", f"v{number + 1}"
        edited = ids[number - 1]
        store.ledger.put(
            "version_plan",
            target,
            {
                "base_version_id": previous,
                "edits": {edited: copy.deepcopy(EXTRACTION)},
                "retry_assets": [],
                "history": [{"asset_id": edited, "reason": "人工对照原图确认"}],
            },
        )
        result = processor.process(target, pictures, {}, "doc")
        assert analyzer.calls == 0
        for index, (_, asset) in enumerate(result):
            assert store.read_image(target, asset) == pictures[index].data
            if index < number:
                assert asset["status"] == "verified"
                assert asset["origin"] == "human_review"
                assert len(asset["history"]) == 1
            elif index < 5:
                assert asset["status"] == "needs_review"
                for key in ("extraction", "issues", "audit", "analyzed_at", "revision"):
                    assert asset[key] == base[asset["id"]][key]
                assert review_issues(asset) == review_issues(base[asset["id"]])
                assert asset["inherited_from_version_id"] == previous
        assert store.ledger.get("version", target)["inherited_images"] == 5
        assert store.ledger.usage_report(target)["call_count"] == 0
    # Re-executing the same immutable job does not duplicate review receipts.
    again = processor.process("v4", pictures, {}, "doc")
    assert analyzer.calls == 0
    assert len(again[2][1]["history"]) == 1


def test_retry_only_calls_the_selected_image_even_if_other_snapshots_failed(revision_workspace):
    store, analyzer, processor, pictures = revision_workspace
    original = store.list_assets("v1")
    failed = original[3]
    failed.update(status="failed", stage="failed", error_code="OCR_PROVIDER_UNAVAILABLE")
    store.ledger.upsert_asset("v1", failed)
    store.ledger.put(
        "version_plan", "v2", {"base_version_id": "v1", "retry_assets": [original[1]["id"]]}
    )
    result = processor.process("v2", pictures, {}, "doc")
    assert analyzer.calls == 1
    assert result[1][1]["status"] == "verified"
    assert result[0][1]["status"] == "needs_review"
    assert result[3][1]["status"] == "failed"
    assert result[3][1]["error_code"] == failed["error_code"]
    assert result[3][1]["issues"] == failed["issues"]


@pytest.mark.parametrize("changed", ["context", "image", "incomplete", "missing"])
def test_unselected_changed_or_incomplete_sources_are_blocked_without_automatic_ocr(
    revision_workspace, changed
):
    store, analyzer, processor, pictures = revision_workspace
    original = store.list_assets("v1")
    if changed == "context":
        pictures[2] = replace(pictures[2], context="different product")
    elif changed == "image":
        pictures[2] = replace(pictures[2], locator=SourceLocator(part="replaced-page.png"))
    elif changed == "incomplete":
        original[2].update(status="queued", stage="extracting")
        store.ledger.upsert_asset("v1", original[2])
    else:
        pictures.append(SourceImage(png(), SourceLocator(part="unexpected-page.png")))
    store.ledger.put(
        "version_plan", "v2", {"base_version_id": "v1", "edits": {original[0]["id"]: EXTRACTION}}
    )
    result = processor.process("v2", pictures, {}, "doc")
    assert analyzer.calls == 0
    blocked = result[-1 if changed == "missing" else 2][1]
    assert blocked["status"] == "needs_review"
    assert blocked["extraction"] is None
    assert blocked["error_code"] == "VISUAL_REVISION_SOURCE_CHANGED"


def test_snapshot_reuse_does_not_call_local_ocr_or_promote_pending_via_global_cache(
    revision_workspace, monkeypatch
):
    store, analyzer, processor, pictures = revision_workspace
    processor.settings = processor.settings.model_copy(update={"ocr_local_check_enabled": True})

    def unexpected(*args, **kwargs):
        raise AssertionError("An inherited snapshot must not be reprocessed")

    monkeypatch.setattr("ragkb.document_processing.visual_processing.read_regions", unexpected)
    monkeypatch.setattr(store.ledger, "cache_get", unexpected)
    store.ledger.put("version_plan", "v2", {"base_version_id": "v1"})
    result = processor.process("v2", pictures, {}, "doc")
    assert analyzer.calls == 0
    assert [a["status"] for _, a in result] == ["needs_review"] * 5 + ["verified"]


def test_inherited_pending_content_stays_out_of_retrieval_and_blocks_publication(tmp_path):
    import io

    from docx import Document
    from ragkb.document_processing.office_parsers import DOCXParser
    from ragkb.document_processing.visual_parser import VisualDocumentParser

    source = tmp_path / "two-images.docx"
    doc = Document()
    doc.add_heading("Independent sections", 1)
    doc.add_picture(io.BytesIO(png()))
    doc.add_picture(io.BytesIO(png()))
    doc.save(source)
    store = VisualAssetStore(LocalFileStorage(tmp_path / "store"))
    analyzer = CountingAnalyzer()
    parser = VisualDocumentParser(DOCXParser(), "docx", analyzer, store, settings())
    parser.parse(source, "base")
    assets = store.list_assets("base")
    assert len(assets) == 2
    assets[1].update(status="needs_review", stage="needs_review", issues=["未经复核"])
    assets[1]["extraction"]["description"] = "UNREVIEWED_SECRET_FACT"
    store.ledger.upsert_asset("base", assets[1])
    analyzer.calls = 0
    store.ledger.put(
        "version_plan",
        "edited",
        {"base_version_id": "base", "edits": {assets[0]["id"]: EXTRACTION}},
    )
    result = parser.parse(source, "edited")
    assert analyzer.calls == 0
    assert "VISUAL_REVIEW_REQUIRED:" + assets[1]["id"] in result.quality_issues
    assert "UNREVIEWED_SECRET_FACT" not in "\n".join(n.original_text for n in result.nodes)
    assert store.get("edited", assets[1]["id"])["status"] == "needs_review"
