"""Regression cases for image placement, atomic table rows, evidence veto and image views."""

from __future__ import annotations

import base64
import io
from dataclasses import replace

import pytest
from docx import Document
from PIL import Image, ImageDraw
from ragkb.adapters.local_storage import LocalFileStorage
from ragkb.adapters.visual_http import VisualAnalyzer
from ragkb.contracts.ports import ParsingDeferred
from ragkb.document_processing.chunking import ChunkingConfig, SemanticChunker, TokenAwareChunker
from ragkb.document_processing.docx_positions import paragraph_positions
from ragkb.document_processing.image_views import TILE_SIZE, image_views, tile_offsets
from ragkb.document_processing.parsers import DOCXParser, ImageParserRoute
from ragkb.document_processing.visual_parser import VisualDocumentParser, _insert_visual_nodes
from ragkb.document_processing.visual_sources import SourceImage, native_images
from ragkb.domain.documents import CanonicalNode, NodeType, SourceLocator
from ragkb.domain.rag import Evidence, EvidencePackage
from ragkb.infrastructure.visual_assets import VisualAssetStore
from ragkb.infrastructure.visual_evidence import VisualEvidenceEnricher
from test_visual_pipeline import EXTRACTION, PASS, Transport, png, settings


def _parser(tmp_path, kind, transport):
    store = VisualAssetStore(LocalFileStorage(tmp_path / "store"))
    return VisualDocumentParser(
        DOCXParser() if kind == "docx" else ImageParserRoute(),
        kind,
        VisualAnalyzer(settings(), transport),
        store,
        settings(),
    )


def test_docx_custom_heading_uses_inherited_outline_and_body_override():
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    custom = document.styles.add_style("产品标题", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = document.styles["Heading 1"]
    document.add_paragraph("Product A", style=custom)
    ordinary = document.add_paragraph("This is ordinary text", style=custom)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    ordinary._p.get_or_add_pPr().append(outline)
    document.add_paragraph("Product B", style=custom)
    document.add_paragraph("B illustration")
    positions = paragraph_positions(document)
    assert [p.section_path for p in positions] == [
        "Product A",
        "Product A",
        "Product B",
        "Product B",
    ]
    assert positions[1].heading_level is None


def test_oversized_image_reports_specific_reason_without_provider_retry():
    transport = Transport()
    analyzer = VisualAnalyzer(settings().model_copy(update={"ocr_max_image_pixels": 1}), transport)
    result = analyzer.analyze(png())
    assert result["status"] == "needs_review"
    assert result["issues"] == ["OCR_IMAGE_PIXELS_LIMIT"]
    assert len(result["audit"]) == 1
    assert transport.calls == []


@pytest.mark.parametrize("strategy", ["structure", "token", "semantic"])
def test_docx_pictures_keep_section_and_context_across_headings_and_blank_paragraphs(
    tmp_path, strategy
):
    path = tmp_path / "two-products.docx"
    document = Document()
    document.add_heading("Product A", 1)
    document.add_paragraph("")
    document.add_paragraph("A-only usage conditions")
    document.add_picture(io.BytesIO(png()))
    document.add_paragraph("Figure 1: Product A diagram", style="Caption")
    document.add_heading("Product B", 1)
    document.add_paragraph("B-only operating voltage")
    document.add_picture(io.BytesIO(png()))
    document.save(path)
    transport = Transport(EXTRACTION, PASS, EXTRACTION, PASS)
    parsed = _parser(tmp_path, "docx", transport).parse(path, "v")
    images = [node for node in parsed.nodes if node.node_type is NodeType.IMAGE]
    assert len(images) == 2
    assert images[0].locator.paragraph == 4
    assert images[0].metadata["section_path"] == "Product A"
    assert images[1].metadata["section_path"] == "Product B"
    assert parsed.nodes.index(images[0]) < next(
        i for i, node in enumerate(parsed.nodes) if node.original_text == "Product B"
    )
    assert "B-only" not in images[0].original_text
    assert "Product B" not in transport.calls[0][2]["messages"][0]["content"]
    config = ChunkingConfig(strategy=strategy)
    chunker = (
        SemanticChunker(lambda _a, _b: 1.0, config=config)
        if strategy == "semantic"
        else TokenAwareChunker(config)
    )
    chunks = chunker.chunk(parsed, tenant_id="test")
    for image in images:
        matching = [
            c
            for c in chunks.chunks
            if c.metadata.get("visual_asset_ids") == image.metadata["visual_asset_ids"]
        ]
        assert matching
        assert all(c.metadata["section_path"] == image.metadata["section_path"] for c in matching)
        wrong = "Product B" if image is images[0] else "Product A"
        assert all(wrong not in c.retrieval_text for c in matching)


def test_unlocated_picture_never_inherits_last_heading():
    title = CanonicalNode(
        "h",
        None,
        NodeType.HEADING,
        "Product B",
        "Product B",
        SourceLocator(page=2),
        {"heading_level": 1},
    )
    picture = CanonicalNode(
        "p",
        None,
        NodeType.IMAGE,
        "Unlocated diagram",
        "Unlocated diagram",
        SourceLocator(part="media/x.png"),
        {"section_path": "root", "visual_asset_ids": ["x"]},
    )
    from ragkb.domain.documents import CanonicalDocument

    document = CanonicalDocument(
        "v", "en", "pdf", tuple(_insert_visual_nodes([title], [picture])), "test", "test", "a" * 64
    )
    result = TokenAwareChunker().chunk(document, tenant_id="test")
    assert result.chunks[0].metadata["section_path"] == "root"
    assert result.chunks[0].metadata["heading"] == ""


def large_table():
    cells = [
        {"row": 0, "column": 0, "rowspan": 2, "colspan": 1, "text": "型号"},
        {"row": 0, "column": 1, "rowspan": 1, "colspan": 2, "text": "电气参数"},
        {"row": 1, "column": 1, "rowspan": 1, "colspan": 1, "text": "工作功率 (W)"},
        {"row": 1, "column": 2, "rowspan": 1, "colspan": 1, "text": "待机功率 (W)"},
    ]
    for row in range(2, 62):
        for column, value in enumerate([f"Orion-{row}", str(200 + row), str(row)]):
            cells.append({"row": row, "column": column, "rowspan": 1, "colspan": 1, "text": value})
    return {
        "title": "设备功率",
        "rows": 62,
        "columns": 3,
        "header_rows": 2,
        "notes": ["仅适用于 220 V 电源"],
        "cells": cells,
    }


def test_provider_table_is_replaced_in_place_without_leaving_unchecked_flattened_copy(tmp_path):
    from ragkb.document_processing.parser_common import canonical_document

    path = tmp_path / "legacy.doc"
    path.write_bytes(b"synthetic provider source")
    location = SourceLocator(page=1, bbox=(10.0, 30.0, 200.0, 300.0))
    nodes = [
        CanonicalNode("a", None, NodeType.HEADING, "Product A", "Product A", SourceLocator(page=1)),
        CanonicalNode("old", None, NodeType.TABLE, "OLD_FLAT_999W", "OLD_FLAT_999W", location),
        CanonicalNode("b", None, NodeType.HEADING, "Product B", "Product B", SourceLocator(page=1)),
    ]

    class ProviderParser:
        def parse(self, source, version):
            return canonical_document(source, version, "doc", "test", nodes)

        def visual_sources(self, _document, _max_bytes):
            return [SourceImage(png(), location, "Figure 1", caption="Figure 1")]

    extraction = {**EXTRACTION, "kind": "mixed", "tables": [large_table()]}
    parser = VisualDocumentParser(
        ProviderParser(),
        "doc",
        VisualAnalyzer(settings(), Transport(extraction, PASS)),
        VisualAssetStore(LocalFileStorage(tmp_path / "store")),
        settings(),
    )
    document = parser.parse(path, "v")
    assert [node.node_type for node in document.nodes] == [
        NodeType.HEADING,
        NodeType.IMAGE,
        NodeType.TABLE,
        NodeType.HEADING,
    ]
    assert all(node.metadata["section_path"] == "Product A" for node in document.nodes[1:3])
    chunks = TokenAwareChunker().chunk(document, tenant_id="test")
    assert all(
        "OLD_FLAT_999W" not in chunk.retrieval_text
        for chunk in (*chunks.chunks, *chunks.parent_chunks)
    )
    visual_chunks = [c for c in chunks.chunks if c.metadata.get("visual_asset_ids")]
    assert visual_chunks and all("Product B" not in c.retrieval_text for c in visual_chunks)


def test_unpositioned_pdf_image_does_not_copy_other_products_from_same_page(tmp_path, monkeypatch):
    from types import SimpleNamespace

    page = SimpleNamespace(
        images=[SimpleNamespace(data=png(), name="x.png")],
        extract_text=lambda: "Product A power 100 W. Product B power 999 W.",
    )
    monkeypatch.setattr(
        "ragkb.document_processing.visual_sources.PdfReader",
        lambda _source: SimpleNamespace(pages=[page]),
    )
    picture = native_images(tmp_path / "two-products.pdf", "pdf", 100000)[0]
    assert picture.locator.page == 1
    assert picture.section_path == "root" and picture.context == ""


@pytest.mark.parametrize("strategy", ["structure", "token", "semantic"])
def test_large_image_table_is_indexed_as_complete_rows_with_headers_units_and_notes(
    tmp_path, strategy
):
    path = tmp_path / "table.png"
    path.write_bytes(png())
    extraction = {
        **EXTRACTION,
        "kind": "table",
        "graphs": [],
        "tables": [large_table()],
        "transcription": "DO_NOT_INDEX_FLATTENED_CELL_COPY",
    }
    parsed = _parser(tmp_path, "image", Transport(extraction, PASS)).parse(path, "tablev")
    assert all(node.node_type is NodeType.TABLE for node in parsed.nodes)
    config = ChunkingConfig(strategy=strategy, target_tokens=220, max_tokens=300, overlap_tokens=20)
    chunker = (
        SemanticChunker(lambda _a, _b: 1.0, config=config)
        if strategy == "semantic"
        else TokenAwareChunker(config)
    )
    chunks = chunker.chunk(parsed, tenant_id="test").chunks
    assert len(chunks) > 1
    covered = []
    for chunk in chunks:
        assert "电气参数 / 工作功率 (W)" in chunk.original_text
        assert "电气参数 / 待机功率 (W)" in chunk.original_text
        assert "仅适用于 220 V 电源" in chunk.original_text
        assert "DO_NOT_INDEX" not in chunk.retrieval_text and "<td" not in chunk.original_text
        assert chunk.token_count <= 300
        first, last = chunk.metadata["visual_table_row_range"]
        covered.extend(range(first, last + 1))
        for row in range(first - 1, last):
            assert f"| Orion-{row} | {200 + row} | {row} |" in chunk.original_text
        assert chunk.metadata["source_spans"][0]["header_row_range"] == [1, 2]
    assert covered == list(range(3, 63))


def test_oversized_single_cell_is_rejected_instead_of_cutting_it(tmp_path):
    path = tmp_path / "table.png"
    path.write_bytes(png())
    table = large_table()
    table["cells"][-1]["text"] = "这是一条不能从中间切断的适用条件" * 100
    extraction = {**EXTRACTION, "kind": "table", "graphs": [], "tables": [table]}
    parsed = _parser(tmp_path, "image", Transport(extraction, PASS)).parse(path, "v")
    with pytest.raises(ParsingDeferred) as result:
        TokenAwareChunker().chunk(parsed, tenant_id="test")
    assert result.value.code == "VISUAL_TABLE_ROW_TOO_LARGE"


def _evidence(asset, identity="E1", chunk="image-chunk", text="Power is 420 W."):
    return Evidence(
        identity,
        chunk,
        "doc",
        "v",
        text,
        {"part": "image.png", "visual_asset_ids": [asset["id"]]},
        0,
        0,
        0,
        1,
        True,
        True,
    )


@pytest.mark.parametrize("status", ["not_relevant", "uncertain", "conflict", "verification_failed"])
def test_rejected_image_cannot_reenter_through_parent_or_supplemental_retrieval(tmp_path, status):
    store = VisualAssetStore(LocalFileStorage(tmp_path))
    asset = store.save_image("v", png(), {"part": "image.png"})
    asset.update(status="verified", extraction=EXTRACTION)
    store.write_manifest("v", [asset])
    replies = (
        [
            {"status": "supported", "text": "Power is 80 W.", "uncertainties": []},
            {**PASS, "text_correct": False, "issues": ["数字模糊"]},
        ]
        if status == "verification_failed"
        else [{"status": status, "text": "", "uncertainties": ["不可确定"]}]
    )
    transport = Transport(*replies)
    session = VisualEvidenceEnricher(store, VisualAnalyzer(settings(), transport), 4).session(
        "Power?"
    )
    image = _evidence(asset)
    parent = replace(
        image,
        evidence_id="E2",
        chunk_id="parent",
        source_role="parent_context",
        text="Other facts. Power is 420 W.",
    )
    plain = replace(
        image,
        evidence_id="E3",
        chunk_id="safe-text",
        text="Warranty is three years.",
        locator={"page": 9},
    )
    initial = session((image, parent, plain))
    assert len(initial) == 1 and initial[0].chunk_id == "safe-text"
    assert initial[0].evidence_id == "E1" and "420" not in initial[0].text
    later_parent = replace(parent, evidence_id="E2", chunk_id="another-parent")
    assert session((*initial, later_parent)) == initial
    assert not transport.responses and len(transport.calls) == len(replies)
    assert "VISUAL_EVIDENCE_EXCLUDED:" + status in session.warnings
    EvidencePackage("run", "tenant", "user", "Power?", 0, "g", "r", "p", "m", 1, initial)


def test_visual_budget_and_questions_are_isolated(tmp_path):
    store = VisualAssetStore(LocalFileStorage(tmp_path))
    assets = [store.save_image("v", png(), {"part": f"{i}.png"}) for i in range(2)]
    for asset in assets:
        asset.update(status="verified", extraction=EXTRACTION)
    store.write_manifest("v", assets)
    transport = Transport(
        {"status": "supported", "text": "Power is 420 W.", "uncertainties": []},
        PASS,
        {"status": "uncertain", "text": "", "uncertainties": ["新问题无法确认"]},
    )
    enrich = VisualEvidenceEnricher(store, VisualAnalyzer(settings(), transport), 1)
    evidence = tuple(_evidence(asset, f"E{i + 1}", f"chunk-{i}") for i, asset in enumerate(assets))
    session = enrich.session("Power?")
    assert len(session(evidence)) == 1
    assert "VISUAL_EVIDENCE_EXCLUDED:budget_exceeded" in session.warnings
    assert enrich.session("Different question?")((evidence[0],)) == ()
    assert len(transport.calls) == 3


def test_evidence_builder_does_not_send_rejected_old_facts_to_selection(tmp_path):
    from ragkb.contracts.rag import EvidenceSelection
    from ragkb.domain.retrieval import SearchHit, SearchResult, SearchSource
    from test_rag_general_repair import provider

    store = VisualAssetStore(LocalFileStorage(tmp_path))
    asset = store.save_image("v", png(), {"part": "image.png"})
    asset.update(status="verified", extraction=EXTRACTION)
    store.write_manifest("v", [asset])

    class Search:
        revision = "test"
        calls = 0

        def search(self, question, context, **kwargs):
            self.calls += 1
            # Supplemental retrieval returns a different parent containing the same rejected image.
            hit = SearchHit(
                f"visual-{self.calls}",
                "doc",
                "v",
                "WRONG_420_W",
                {"page": 1, "visual_asset_ids": [asset["id"]]},
                1,
                1,
                ("dense",),
                display_text="WRONG_420_W",
                retrieval_text="WRONG_420_W",
            )
            safe = SearchSource(
                "safe",
                "doc",
                "v",
                "Warranty 3 years",
                "Warranty 3 years",
                {"page": 2},
                0,
                0,
                1,
                True,
            )
            return SearchResult((hit,), 1, review_sources=(safe,))

        def expand_parents(self, sources, context):
            return ()

    class Selector:
        revision = "test"

        def select(self, question, evidence):
            assert all("WRONG_420" not in item.text for item in evidence)
            return EvidenceSelection(
                tuple(item.evidence_id for item in evidence), "partial", ("Check another section",)
            )

    transport = Transport(
        {"status": "conflict", "text": "", "uncertainties": ["Old number differs"]}
    )
    search = Search()
    builder = provider(search, Selector())
    builder.visual_enricher = VisualEvidenceEnricher(
        store, VisualAnalyzer(settings(), transport), 4
    )
    package = builder.build_package("Power and warranty?", "tenant", "user")
    assert search.calls == 2 and len(transport.calls) == 1
    assert len(package.evidence) == 1 and package.evidence[0].chunk_id == "safe"
    assert "VISUAL_EVIDENCE_EXCLUDED:conflict" in package.retrieval_warnings


def _prepared(data):
    return image_views(data, max_bytes=20 * 1024 * 1024, max_pixels=40_000_000)


def _view(parts, index=0):
    return Image.open(
        io.BytesIO(base64.b64decode(parts[index * 2 + 1]["image_url"]["url"].split(",", 1)[1]))
    )


@pytest.mark.parametrize("palette", [False, True])
def test_transparent_background_is_white_and_black_foreground_remains_visible(palette):
    picture = Image.new("RGBA", (30, 30), (0, 0, 0, 0))
    ImageDraw.Draw(picture).rectangle((10, 10, 20, 20), fill=(0, 0, 0, 255))
    if palette:
        picture = picture.quantize()
    data = io.BytesIO()
    picture.save(data, format="PNG")
    normalized = _view(_prepared(data.getvalue()))
    assert normalized.getpixel((0, 0)) == (255, 255, 255)
    assert normalized.getpixel((15, 15)) == (0, 0, 0)


def test_exif_orientation_is_applied_before_preparing_views():
    picture = Image.new("RGB", (30, 60), "red")
    exif = Image.Exif()
    exif[274] = 6
    data = io.BytesIO()
    picture.save(data, format="JPEG", exif=exif)
    normalized = _view(_prepared(data.getvalue()))
    assert normalized.size == (60, 30)
    assert normalized.getexif().get(274) in {None, 1}


@pytest.mark.parametrize("size", [(200, 12000), (12000, 200)])
def test_long_picture_has_continuous_overlapping_original_resolution_tiles(size):
    picture = Image.new("RGB", size, "white")
    drawing = ImageDraw.Draw(picture)
    drawing.rectangle(
        (size[0] // 2, size[1] // 2, size[0] // 2 + 20, size[1] // 2 + 20), fill="black"
    )
    data = io.BytesIO()
    picture.save(data, format="PNG")
    views = _prepared(data.getvalue())
    assert len(views) // 2 > 5
    positions = tile_offsets(max(size))
    assert positions[0] == 0 and positions[-1] + TILE_SIZE == max(size)
    assert all(
        right < left + TILE_SIZE for left, right in zip(positions, positions[1:], strict=False)
    )
    crops = [_view(views, index) for index in range(1, len(views) // 2)]
    assert all(max(crop.size) <= TILE_SIZE for crop in crops)
    assert any(crop.getextrema()[0][0] == 0 for crop in crops)
