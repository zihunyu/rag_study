"""Behavioral coverage for the three-stage visual workspace, using isolated sources."""

from __future__ import annotations

import copy
import io
import threading
import time
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from pptx import Presentation
from pptx.util import Inches
from pydantic import SecretStr
from ragkb.adapters.local_storage import LocalFileStorage
from ragkb.adapters.visual_http import VisualAnalyzer
from ragkb.application.reading_scope import ReadingOptions, reading_scope
from ragkb.contracts.ports import ParsingDeferred
from ragkb.document_processing.local_visual_check import check_extraction
from ragkb.document_processing.office_parsers import DOCXParser, PPTXParser, SpreadsheetParser
from ragkb.document_processing.visual_coverage import inventory, render_fallback
from ragkb.document_processing.visual_parser import VisualDocumentParser
from ragkb.document_processing.visual_processing import VisualProcessor
from ragkb.document_processing.visual_sources import SourceImage, native_images
from ragkb.domain.documents import SourceLocator
from ragkb.domain.errors import ProviderUnavailable
from ragkb.domain.visual_comparison import compare_readings
from ragkb.domain.visuals import VisualExtraction
from ragkb.infrastructure.visual_assets import VisualAssetStore
from test_visual_pipeline import EXTRACTION, PASS, Transport, png, settings
from test_workspace_redesign import publish, upload, workspace  # noqa: F401


def test_ppt_and_excel_anchors_keep_shared_images_on_their_actual_sources(tmp_path):
    presentation = Presentation()
    for title in ("Product A", "Product B"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_picture(io.BytesIO(png()), Inches(1), Inches(2))
    ppt = tmp_path / "anchors.pptx"
    presentation.save(ppt)
    images = native_images(ppt, "pptx", 1_000_000)
    assert [i.locator.slide for i in images] == [1, 2]
    assert "Product A" in images[0].section_path and "Product B" not in images[0].context
    text = PPTXParser().parse(ppt, "v")
    assert text.nodes[0].metadata["section_path"] == images[0].section_path
    workbook = Workbook()
    for title, cell in (("A", "C5"), ("B", "F12")):
        sheet = workbook.create_sheet(title)
        sheet["A1"] = title
        sheet.add_image(SheetImage(io.BytesIO(png())), cell)
    book = tmp_path / "anchors.xlsx"
    workbook.save(book)
    images = native_images(book, "xlsx", 1_000_000)
    assert [(i.locator.sheet, i.locator.cell_range) for i in images] == [("A", "C5"), ("B", "F12")]
    assert {n.metadata["section_path"] for n in SpreadsheetParser().parse(book, "v").nodes} == {
        "A",
        "B",
    }


def test_external_images_are_reported_and_never_fetched_or_rendered(tmp_path):
    source = tmp_path / "external.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr(
            "word/_rels/document.xml.rels",
            "<Relationships><Relationship "
            'Id="r1" Type="drawing/image" TargetMode="External" Target="https://invalid.example/x"/>'
            "</Relationships>",
        )
    report = inventory(source, "docx")
    assert report["objects"][0]["kind"] == "external_image"
    images, report = render_fallback(source, "docx", report, settings(), tmp_path / "temp")
    assert not images and report["render_error"] == "VISUAL_EXTERNAL_IMAGE_REQUIRES_EMBEDDING"
    assert "https://" not in str(report)


@pytest.mark.skipif(
    not Path("C:/Program Files/LibreOffice/program/soffice.com").is_file(),
    reason="real page rendering requires the installed LibreOffice runtime",
)
def test_real_office_render_and_pdf_vector_inventory_report_page_limit(tmp_path):
    source = tmp_path / "shapes.pptx"
    presentation = Presentation()
    for title in ("Architecture A", "Architecture B"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        from pptx.enum.shapes import MSO_SHAPE

        slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(2), Inches(1)
        ).text = "Gateway"
    presentation.save(source)
    report = inventory(source, "pptx")
    assert report["needs_render"]
    images, result = render_fallback(
        source,
        "pptx",
        report,
        settings().model_copy(update={"ocr_render_max_pages": 1}),
        tmp_path / "render",
    )
    assert len(images) == 1 and images[0].locator.page == 1
    assert result["missing_pages"] == [2]
    assert any(o["state"] == "unprocessed" for o in result["objects"])
    from PIL import Image

    with Image.open(io.BytesIO(images[0].data)) as image:
        assert image.width > 1000 and image.getextrema() != ((255, 255),) * 3


def test_local_ocr_does_not_confirm_23_using_323_or_guess_repeated_locations():
    extraction = VisualExtraction.model_validate(
        {**EXTRACTION, "kind": "text", "transcription": "23 W", "description": "", "graphs": []}
    )
    reading = {
        "engine": "independent",
        "regions": [
            {"id": "r1", "text": "323 W", "score": 0.99},
            {"id": "r2", "text": "Gateway", "score": 0.99},
            {"id": "r3", "text": "Gateway", "score": 0.99},
        ],
    }
    assert check_extraction(extraction, reading)["status"] == "disagreement"
    result = check_extraction(VisualExtraction.model_validate(EXTRACTION), reading)
    assert next(t for t in result["targets"] if t["node_id"] == "a")["region_ids"] == []


@pytest.mark.parametrize("mutation", ["direction", "style", "group"])
def test_independent_readings_detect_arrow_and_containment_changes(mutation):
    other = copy.deepcopy(EXTRACTION)
    if mutation == "group":
        other["graphs"][0]["nodes"][1]["group"] = None
    else:
        other["graphs"][0]["edges"][0][mutation] = "forward" if mutation == "direction" else "solid"
    assert compare_readings(
        VisualExtraction.model_validate(EXTRACTION), VisualExtraction.model_validate(other)
    )


def test_blind_model_verification_uses_independent_endpoint_and_checks_prose():
    config = settings().model_copy(
        update={
            "ocr_verify_enabled": True,
            "ocr_verify_base_url": "https://independent.example/v1",
            "ocr_verify_model": "independent",
            "ocr_verify_api_key": SecretStr("independent-test-secret"),
        }
    )
    transport = Transport(EXTRACTION, EXTRACTION, PASS)
    result = VisualAnalyzer(config, transport).analyze(png())
    assert result["status"] == "verified"
    assert "independent.example" in transport.calls[1][0]
    assert "候选转录" not in transport.calls[1][2]["messages"][0]["content"]
    assert len(transport.calls) == 3


class CountingAnalyzer:
    revision = "counting-v1"

    def __init__(self):
        self.calls = 0
        self.active = 0
        self.maximum = 0
        self.lock = threading.Lock()
        self.fail = False

    def analyze(self, data, context="", *, progress=None):
        with self.lock:
            self.calls += 1
            self.active += 1
            self.maximum = max(self.maximum, self.active)
        try:
            if progress:
                progress("extracting")
            time.sleep(0.03)
            if self.fail and context == "bad":
                raise ProviderUnavailable("SYNTHETIC_RETRY")
            return {
                "status": "verified",
                "extraction": copy.deepcopy(EXTRACTION),
                "issues": [],
                "revision": self.revision,
                "analyzed_at": time.time(),
            }
        finally:
            with self.lock:
                self.active -= 1


def test_parallel_processing_resumes_only_failed_images_and_preserves_all_stages(tmp_path):
    store = VisualAssetStore(LocalFileStorage(tmp_path / "store"))
    analyzer = CountingAnalyzer()
    config = settings().model_copy(update={"ocr_max_concurrency": 2})
    processor = VisualProcessor(store, analyzer, config)
    pictures = [
        SourceImage(png(), SourceLocator(part=f"img{i}.png"), "bad" if i == 2 else "ok")
        for i in range(5)
    ]
    analyzer.fail = True
    with pytest.raises(ProviderUnavailable):
        processor.process("v", pictures, {}, "doc")
    assert analyzer.maximum == 2
    assert [a["stage"] for a in store.list_assets("v")].count("completed") == 4
    analyzer.fail = False
    processor.process("v", pictures, {}, "doc")
    assert analyzer.calls == 6
    assert all(a["stage"] == "completed" for a in store.list_assets("v"))


def test_cross_version_cache_and_human_changes_do_not_cross_scopes(tmp_path):
    store = VisualAssetStore(LocalFileStorage(tmp_path / "store"))
    analyzer = CountingAnalyzer()
    config = settings().model_copy(update={"ocr_cross_version_cache_enabled": True})
    processor = VisualProcessor(store, analyzer, config, scope="tenant-a")
    pictures = [SourceImage(png(), SourceLocator(part="original.png"))]
    a = processor.process("v1", pictures, {}, "doc")[0][1]
    b = processor.process("v2", pictures, {}, "doc")[0][1]
    assert analyzer.calls == 1 and b["cache_hit"]
    assert b["storage_key"].startswith(store.prefix("v2"))
    changed = copy.deepcopy(EXTRACTION)
    changed["title"] = "人工修正"
    store.ledger.put(
        "version_plan",
        "v3",
        {
            "base_version_id": "v2",
            "edits": {a["id"]: changed},
            "history": [{"reason": "看原图修改"}],
        },
    )
    processor.process("v3", pictures, {}, "doc")
    store.ledger.put("version_plan", "v4", {"base_version_id": "v3"})
    revised = processor.process("v4", pictures, {}, "doc")[0][1]
    assert revised["origin"] == "human_review" and revised["history"]
    assert revised["history"][0]["asset_id"] == a["id"]
    assert revised["history"][0]["legacy"]
    assert revised["history"][0]["binding_basis"] == "single_edited_asset"
    fresh = processor.process("v5", pictures, {}, "doc")[0][1]
    assert fresh["extraction"]["title"] != "人工修正"
    VisualProcessor(store, analyzer, config, scope="tenant-b").process("vb", pictures, {}, "doc")
    assert analyzer.calls == 2


def test_partial_publication_removes_whole_chapter_and_rejects_empty_result(tmp_path):
    path = tmp_path / "partial.docx"
    doc = Document()
    doc.add_heading("Product A", 1)
    doc.add_paragraph("A-only forbidden facts")
    doc.add_picture(io.BytesIO(png()))
    doc.add_heading("Product B", 1)
    doc.add_paragraph("B independent information")
    doc.save(path)
    store = VisualAssetStore(LocalFileStorage(tmp_path / "store"))
    parser = VisualDocumentParser(DOCXParser(), "docx", CountingAnalyzer(), store, settings())
    store.ledger.put("version_plan", "partial", {"exclude_sections": ["Product A"]})
    result = parser.parse(path, "partial")
    assert "A-only" not in str(result.nodes)
    assert "B independent" in str(result.nodes)
    assert "PARTIAL_PUBLICATION_CONTENT_EXCLUDED" in result.quality_issues
    store.ledger.put("version_plan", "empty", {"exclude_sections": ["Product A", "Product B"]})
    with pytest.raises(ParsingDeferred, match="No independently publishable"):
        parser.parse(path, "empty")


def test_scope_options_are_durable_and_part_of_idempotent_submission(workspace):  # noqa: F811
    client, runtime, service, space = workspace
    from ragkb.domain.auth import RequestPrincipal
    from ragkb.domain.uploads import IdempotencyConflictError

    subject = RequestPrincipal(runtime.tenant_id, "local-user", ("admin",), (), "local_single_user")
    repository = service.repository
    conversation = repository.create(space, "测试", subject, "scope-conversation")
    reading = {"mode": "overview", "document_ids": ["doc"]}
    turn = repository.submit(conversation["id"], subject, "全文总结", "scope-turn", reading)
    assert repository.turn(turn["id"])["reading"] == reading
    assert (
        repository.submit(conversation["id"], subject, "全文总结", "scope-turn", reading)["id"]
        == turn["id"]
    )
    with pytest.raises(IdempotencyConflictError):
        repository.submit(conversation["id"], subject, "全文总结", "scope-turn", {"mode": "fact"})


def test_overview_reads_last_chapter_and_never_reuses_revoked_content(workspace):  # noqa: F811
    client, runtime, service, space = workspace
    body = "\n".join(
        f"# Chapter {i}\nModel X{i} has a unique capacity of {100 + i} litres.\n" for i in range(12)
    )
    done = upload(client, runtime, space, "chapters.md", body)
    publish(client, done["document_version_id"])
    provider = runtime.qa_service.evidence_provider
    with reading_scope(ReadingOptions(mode="overview", document_ids=(done["document_id"],))):
        package = provider.build_package(
            "Summarize the whole document", runtime.tenant_id, "local-user", space_id=space
        )
    assert "111 litres" in " ".join(e.text for e in package.evidence)
    assert package.coverage_report["read_sections"] >= 12
    assert all(e.document_version_id == done["document_version_id"] for e in package.evidence)
    revoked = client.post(
        f"/api/documents/{done['document_id']}:revoke",
        headers={"Idempotency-Key": "overview-revoke"},
    )
    assert revoked.status_code == 200
    with reading_scope(ReadingOptions(mode="overview", document_ids=(done["document_id"],))):
        after = provider.build_package(
            "Summarize the whole document", runtime.tenant_id, "local-user", space_id=space
        )
    assert not after.evidence
